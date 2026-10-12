"""
Build the Excel workbook excel/NHS_AE_Analysis.xlsx from the cleaned data.

Every summary figure in the workbook is an Excel formula (SUMIFS, AVERAGEIFS,
INDEX/MATCH, COUNTIFS) over the Data sheet, so the workbook recalculates if
the data is updated.

Run from the project root:  python src/build_excel.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
CLEAN = ROOT / "data" / "clean"
OUT = ROOT / "excel" / "NHS_AE_Analysis.xlsx"

FONT = "Arial"
NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COUNT = "#,##0"
PCT = "0.0%"
CHANGE = '+0.0%;-0.0%;0.0%'
POINTS = '+0.0" pts";-0.0" pts";0.0" pts"'


def font(**kw):
    return Font(name=FONT, **kw)


def style_header(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def load_data() -> pd.DataFrame:
    month = pd.read_csv(CLEAN / "dim_month.csv", parse_dates=["month_start"])
    act = pd.read_csv(CLEAN / "fact_activity.csv")
    perf = pd.read_csv(CLEAN / "fact_performance.csv")
    df = month.merge(act, on="month_key").merge(perf, on="month_key", how="left")
    # 4-hour data starts in November 2010; keep only months that have it
    return df[df["within4_total"].notna()].reset_index(drop=True)


# ---------------------------------------------------------------- Data sheet
DATA_COLS = [
    ("Month", "month_start", "mmm yyyy"),
    ("Month label", "month_label", None),
    ("Financial year", "financial_year", None),
    ("Month name", "month_name", None),
    ("Season", "season", None),
    ("Total attendances", "att_total", COUNT),
    ("Type 1 attendances", "att_type1", COUNT),
    ("Type 3 attendances", "att_type3", COUNT),
    ("Emergency admissions", "adm_total", COUNT),
    ("Admitted via Type 1", "adm_via_type1", COUNT),
    ("Seen within 4h (all)", "within4_total", COUNT),
    ("Over 4h (all)", "over4_total", COUNT),
    ("Seen within 4h (Type 1)", "within4_type1", COUNT),
    ("Over 4h (Type 1)", "over4_type1", COUNT),
    ("12h waits after decision to admit", "dta_over_12hrs", COUNT),
    ("Flag: CRS field test", "flag_crs_field_test", "0"),
    ("Flag: estimated from weekly", "flag_estimated_from_weekly", "0"),
]


def build_data_sheet(wb, df):
    ws = wb.create_sheet("Data")
    headers = [h for h, _, _ in DATA_COLS] + [
        "% within 4h (all)", "% within 4h (Type 1)", "Type 1 admission rate", "In seasonality window"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, 1, len(headers))
    ws.row_dimensions[1].height = 45

    n = len(df)
    for i, rec in enumerate(df.itertuples(index=False), start=2):
        for c, (_, field, fmt) in enumerate(DATA_COLS, 1):
            val = getattr(rec, field)
            if field == "month_start":
                val = val.to_pydatetime()
            elif pd.isna(val):
                val = None
            elif fmt in (COUNT, "0"):
                val = int(val)
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = font()
            if fmt:
                cell.number_format = fmt
        # Calculated columns (R to U)
        ws.cell(row=i, column=18, value=f"=K{i}/(K{i}+L{i})").number_format = PCT
        ws.cell(row=i, column=19, value=f"=M{i}/(M{i}+N{i})").number_format = PCT
        ws.cell(row=i, column=20, value=f"=J{i}/G{i}").number_format = PCT
        ws.cell(row=i, column=21,
                value=f"=IF(AND(A{i}>=Seasonality!$C$4,A{i}<=Seasonality!$C$5),1,0)")
        for c in range(18, 22):
            ws.cell(row=i, column=c).font = font()

    last = n + 1
    table = Table(displayName="tblMonthly", ref=f"A1:{get_column_letter(len(headers))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "C2"
    set_widths(ws, {get_column_letter(c): 13 for c in range(1, len(headers) + 1)})
    set_widths(ws, {"B": 11, "C": 11, "D": 8, "E": 9})
    ws["A1"].comment = Comment(
        "Source: NHS England, Monthly A&E Time Series September 2026 (published 8 Oct 2026). "
        "Months before June 2015 were estimated by NHS England from weekly data.", "Somy")
    return last


# ----------------------------------------------------------- FY Summary sheet
def build_fy_sheet(wb, df, last):
    ws = wb.create_sheet("FY Summary")
    ws["A1"] = "A&E demand and 4-hour performance by NHS financial year"
    ws["A1"].font = font(bold=True, size=14, color=NAVY)
    ws["A2"] = ("All figures are formulas over the Data sheet (SUMIFS / COUNTIFS). "
                "Years with fewer than 12 months of data are marked as partial.")
    ws["A2"].font = font(italic=True, color="595959")

    headers = ["Financial year", "Months of data", "Total attendances", "Change vs previous year",
               "Type 1 attendances", "Emergency admissions", "% within 4h (all)",
               "% within 4h (Type 1)", "Change in Type 1 % (points)", "12h waits for a bed",
               "Status"]
    hr = 4
    for c, h in enumerate(headers, 1):
        ws.cell(row=hr, column=c, value=h)
    style_header(ws, hr, 1, len(headers))
    ws.row_dimensions[hr].height = 45

    years = sorted(df["financial_year"].unique())
    rng = lambda col: f"Data!${col}$2:${col}${last}"
    for r, fy in enumerate(years, start=hr + 1):
        ws.cell(row=r, column=1, value=fy)
        ws.cell(row=r, column=2, value=f'=COUNTIFS({rng("C")},$A{r})')
        ws.cell(row=r, column=3, value=f'=SUMIFS({rng("F")},{rng("C")},$A{r})')
        ws.cell(row=r, column=4, value=(f'=IF(AND(B{r}=12,B{r-1}=12),C{r}/C{r-1}-1,"")'
                                        if r > hr + 1 else '=""'))
        ws.cell(row=r, column=5, value=f'=SUMIFS({rng("G")},{rng("C")},$A{r})')
        ws.cell(row=r, column=6, value=f'=SUMIFS({rng("I")},{rng("C")},$A{r})')
        ws.cell(row=r, column=7, value=(f'=SUMIFS({rng("K")},{rng("C")},$A{r})/'
                                        f'(SUMIFS({rng("K")},{rng("C")},$A{r})+SUMIFS({rng("L")},{rng("C")},$A{r}))'))
        ws.cell(row=r, column=8, value=(f'=SUMIFS({rng("M")},{rng("C")},$A{r})/'
                                        f'(SUMIFS({rng("M")},{rng("C")},$A{r})+SUMIFS({rng("N")},{rng("C")},$A{r}))'))
        ws.cell(row=r, column=9, value=(f'=IF(AND(B{r}=12,B{r-1}=12),(H{r}-H{r-1})*100,"")'
                                        if r > hr + 1 else '=""'))
        ws.cell(row=r, column=10, value=f'=SUMIFS({rng("O")},{rng("C")},$A{r})')
        ws.cell(row=r, column=11, value=f'=IF(B{r}<12,"Partial year","Complete")')
        fmts = [None, "0", COUNT, CHANGE, COUNT, COUNT, PCT, PCT, '+0.0;-0.0;0.0', COUNT, None]
        for c, f in enumerate(fmts, 1):
            cell = ws.cell(row=r, column=c)
            cell.font = font()
            cell.border = BOX
            if f:
                cell.number_format = f
    end = hr + len(years)

    # Highlight years below the 95% standard
    ws["M4"] = "4-hour standard"
    ws["M4"].font = font(bold=True)
    ws["N4"] = 0.95
    ws["N4"].number_format = "0%"
    ws["N4"].font = font(color="0000FF")
    ws["N4"].fill = INPUT_FILL
    ws["N4"].comment = Comment("NHS constitutional standard: 95% of patients admitted, "
                               "transferred or discharged within 4 hours. Taken from the "
                               "source spreadsheet.", "Somy")
    ws.conditional_formatting.add(
        f"H{hr+1}:H{end}", CellIsRule(operator="lessThan", formula=["$N$4"],
                                      font=Font(name=FONT, color="C00000", bold=True)))
    ws["M5"] = "Type 1 % below the standard is shown in red."
    ws["M5"].font = font(italic=True, color="595959")

    set_widths(ws, {"A": 14, "B": 9, "C": 15, "D": 13, "E": 15, "F": 15, "G": 12, "H": 12,
                    "I": 13, "J": 13, "K": 13, "L": 3, "M": 16, "N": 8})
    ws.freeze_panes = "B5"

    # Chart 1: Type 1 and overall 4-hour performance
    line = LineChart()
    line.title = "% of patients seen within 4 hours"
    line.y_axis.title = "% within 4 hours"
    line.y_axis.number_format = "0%"
    line.y_axis.scaling.min = 0.5
    line.y_axis.scaling.max = 1.0
    line.height, line.width = 9, 22
    line.add_data(Reference(ws, min_col=7, max_col=8, min_row=hr, max_row=end), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=hr + 1, max_row=end))
    for series in line.series:
        series.smooth = False
    line.legend.position = "b"
    ws.add_chart(line, f"A{end + 3}")

    # Chart 2: attendances
    bar = BarChart()
    bar.title = "Total A&E attendances per financial year"
    bar.y_axis.number_format = "#,##0,,\"m\""
    bar.height, bar.width = 9, 22
    bar.legend = None
    bar.add_data(Reference(ws, min_col=3, min_row=hr, max_row=end), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=hr + 1, max_row=end))
    ws.add_chart(bar, f"A{end + 22}")
    return end


# ---------------------------------------------------------- Month Lookup sheet
def build_lookup_sheet(wb, df, last):
    ws = wb.create_sheet("Month Lookup")
    ws["A1"] = "Month lookup: pick any month and compare with a year earlier"
    ws["A1"].font = font(bold=True, size=14, color=NAVY)
    ws["A2"] = "Change the yellow cell (drop-down). Everything else is INDEX/MATCH over the Data sheet."
    ws["A2"].font = font(italic=True, color="595959")

    ws["A4"] = "Selected month"
    ws["A4"].font = font(bold=True)
    ws["B4"] = df["month_label"].iloc[-1]
    ws["B4"].fill = INPUT_FILL
    ws["B4"].font = font(bold=True, color="0000FF")
    ws["B4"].border = BOX
    dv = DataValidation(type="list", formula1=f"=Data!$B$2:$B${last}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")

    ws["A5"] = "Same month last year"
    ws["A5"].font = font(bold=True)
    ws["D4"] = "Row of selected month"
    ws["E4"] = f"=MATCH($B$4,Data!$B$2:$B${last},0)"
    ws["D5"] = "Row of last year"
    ws["E5"] = (f"=IFERROR(MATCH(EDATE(INDEX(Data!$A$2:$A${last},E4),-12),"
                f"Data!$A$2:$A${last},0),NA())")
    ws["B5"] = f'=IFERROR(INDEX(Data!$B$2:$B${last},E5),"Not available")'
    for ref in ("D4", "D5", "E4", "E5"):
        ws[ref].font = font(color="808080", size=9)

    headers = ["Measure", "Selected month", "Same month last year", "Change"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=7, column=c, value=h)
    style_header(ws, 7, 1, 4)

    measures = [  # label, Data column, format, change type
        ("Total attendances", "F", COUNT, "pct"),
        ("Type 1 attendances", "G", COUNT, "pct"),
        ("Emergency admissions", "I", COUNT, "pct"),
        ("% within 4h (all)", "R", PCT, "points"),
        ("% within 4h (Type 1)", "S", PCT, "points"),
        ("Type 1 admission rate", "T", PCT, "points"),
        ("12h waits after decision to admit", "O", COUNT, "pct"),
    ]
    for r, (label, col, fmt, kind) in enumerate(measures, start=8):
        ws.cell(row=r, column=1, value=label).font = font()
        ws.cell(row=r, column=2, value=f"=INDEX(Data!${col}$2:${col}${last},$E$4)").number_format = fmt
        ws.cell(row=r, column=3, value=f'=IFERROR(INDEX(Data!${col}$2:${col}${last},$E$5),"")').number_format = fmt
        if kind == "pct":
            ws.cell(row=r, column=4, value=f'=IFERROR(B{r}/C{r}-1,"")').number_format = CHANGE
        else:
            ws.cell(row=r, column=4, value=f'=IFERROR((B{r}-C{r})*100,"")').number_format = POINTS
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = BOX
            ws.cell(row=r, column=c).font = font()

    ws["A16"] = "Data caveat for selected month"
    ws["A16"].font = font(bold=True)
    ws["B16"] = (f'=IF(INDEX(Data!$P$2:$P${last},$E$4)=1,"CRS field test: 14 trusts missing from 4-hour data",'
                 f'IF(INDEX(Data!$Q$2:$Q${last},$E$4)=1,"Estimated by NHS England from weekly data","None"))')
    ws["B16"].font = font(color="C00000")
    set_widths(ws, {"A": 32, "B": 18, "C": 20, "D": 14, "E": 8})


# ----------------------------------------------------------- Seasonality sheet
def build_season_sheet(wb, last):
    ws = wb.create_sheet("Seasonality")
    ws["A1"] = "Seasonal pattern: average by calendar month"
    ws["A1"].font = font(bold=True, size=14, color=NAVY)
    ws["A2"] = "AVERAGEIFS over the months inside the window below. Change the yellow dates to test other periods."
    ws["A2"].font = font(italic=True, color="595959")

    ws["B4"], ws["B5"] = "Window start", "Window end"
    ws["C4"] = pd.Timestamp("2022-04-01").to_pydatetime()
    ws["C5"] = pd.Timestamp("2026-03-01").to_pydatetime()
    for ref in ("C4", "C5"):
        ws[ref].number_format = "mmm yyyy"
        ws[ref].fill = INPUT_FILL
        ws[ref].font = font(color="0000FF")
        ws[ref].border = BOX
    ws["D4"] = "Default: the four most recent complete financial years (2022-23 to 2025-26)."
    ws["D4"].font = font(italic=True, color="595959")
    ws["B6"] = "Months in window"
    ws["C6"] = f"=SUM(Data!$U$2:$U${last})"

    headers = ["Month", "Avg attendances", "Avg % within 4h (Type 1)", "Avg 12h waits for a bed"]
    for c, h in enumerate(headers, 2):
        ws.cell(row=8, column=c, value=h)
    style_header(ws, 8, 2, 5)
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    for r, m in enumerate(months, start=9):
        ws.cell(row=r, column=2, value=m)
        crit = f'Data!$D$2:$D${last},$B{r},Data!$U$2:$U${last},1'
        ws.cell(row=r, column=3, value=f"=AVERAGEIFS(Data!$F$2:$F${last},{crit})").number_format = COUNT
        ws.cell(row=r, column=4, value=f"=AVERAGEIFS(Data!$S$2:$S${last},{crit})").number_format = PCT
        ws.cell(row=r, column=5, value=f"=AVERAGEIFS(Data!$O$2:$O${last},{crit})").number_format = COUNT
        for c in range(2, 6):
            ws.cell(row=r, column=c).border = BOX
            ws.cell(row=r, column=c).font = font()
    set_widths(ws, {"A": 3, "B": 16, "C": 16, "D": 16, "E": 16})

    chart = BarChart()
    chart.title = "Average % within 4 hours (Type 1) by month"
    chart.y_axis.number_format = "0%"
    chart.y_axis.scaling.min = 0.5
    chart.legend = None
    chart.height, chart.width = 8, 18
    chart.add_data(Reference(ws, min_col=4, min_row=8, max_row=20), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=9, max_row=20))
    ws.add_chart(chart, "G4")


# ---------------------------------------------------------------- About sheet
def build_about_sheet(ws, df):
    ws.title = "About"
    ws["A1"] = "NHS England A&E Performance Analysis"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    rows = [
        ("Purpose", "Track A&E demand, the 4-hour standard and long waits for a bed in England, "
                    "from November 2010 to the latest month."),
        ("Source", "NHS England, Monthly A&E Time Series September 2026, published 8 October 2026. "
                   "https://www.england.nhs.uk/statistics/statistical-work-areas/ae-waiting-times-and-activity/"),
        ("Coverage", f"{df['month_label'].iloc[0]} to {df['month_label'].iloc[-1]} "
                     f"({len(df)} months), England, all NHS and independent sector providers."),
        ("Sheets", "Data: one row per month, with calculated % columns.  "
                   "FY Summary: totals and % by financial year, with charts.  "
                   "Month Lookup: choose a month and compare with a year earlier.  "
                   "Seasonality: average by calendar month over a chosen window."),
        ("Yellow cells", "Inputs you can change. Blue text = typed-in value; black = formula."),
        ("Type 1", "Major emergency departments (consultant-led, 24-hour). "
                   "Type 3 = urgent treatment centres and minor injury units."),
        ("Caveat 1", "Nov 2010 to May 2015 were estimated by NHS England from weekly returns."),
        ("Caveat 2", "May 2019 to May 2023: 14 trusts testing new standards (CRS) are missing from the "
                     "4-hour figures, so performance in this period is not fully comparable."),
        ("Caveat 3", "From Aug 2020 attendances include booked appointments."),
        ("Caveat 4", "Aug 2022 to Feb 2023: a cyber-attack caused a shortfall in Type 3 attendances."),
        ("Licence", "Contains public sector information licensed under the Open Government Licence v3.0."),
    ]
    for r, (k, v) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=k).font = font(bold=True)
        cell = ws.cell(row=r, column=2, value=v)
        cell.font = font()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        ws.cell(row=r, column=1).fill = NOTE_FILL
    set_widths(ws, {"A": 16, "B": 110})


def main():
    df = load_data()
    wb = Workbook()
    build_about_sheet(wb.active, df)
    last = build_data_sheet(wb, df)
    build_fy_sheet(wb, df, last)
    build_lookup_sheet(wb, df, last)
    build_season_sheet(wb, last)
    # Order: About, FY Summary, Month Lookup, Seasonality, Data
    wb._sheets = [wb["About"], wb["FY Summary"], wb["Month Lookup"], wb["Seasonality"], wb["Data"]]
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
